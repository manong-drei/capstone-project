from pathlib import Path
import json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parent.parent
SHOTS=ROOT/'screenshots'
DOC=ROOT/'EKalusugan_User_Manual.docx'
ps=json.loads((ROOT/'work/patient_staff.json').read_text(encoding='utf-8'))
da=json.loads((ROOT/'work/doctor_admin.json').read_text(encoding='utf-8'))
p={s['id']:s for s in ps['sections']}
d=da['doctor']['sections']; a=da['admin']['sections']
p['staff-account-recovery']['notes']=[x for x in p['staff-account-recovery']['notes'] if not x.startswith('Do not include')]
d[2]['notes']=[x.replace('do not describe it as an unlimited allowance','set the intended allowance before accepting walk-ins') for x in d[2]['notes']]
a[10]['notes']=[x for x in a[10]['notes'] if not x.startswith('No password value')]
a[11]['notes']=[x.replace('The result disappears at expiry; it should not be captured in published screenshots.','The result disappears at expiry. Hand it directly to the verified patient and keep it private.') for x in a[11]['notes']]
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.67); sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.65)
sec.header_distance=Inches(.25); sec.footer_distance=Inches(.28)
W=6.97
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Bullet','List Number']:
 st=styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Arial')
 st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(5)
styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',31),('Subtitle',14),('Heading 1',20),('Heading 2',14),('Heading 3',11.5)]:
 styles[name].font.size=Pt(size)
 styles[name].paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 0)
 styles[name].paragraph_format.space_after=Pt(7)
 if name.startswith('Heading'): styles[name].paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(9); styles['Caption'].font.italic=False
styles['Caption'].paragraph_format.space_after=Pt(6)
styles['Caption'].paragraph_format.keep_with_next=True
for name in ['List Bullet','List Number']:
 styles[name].font.size=Pt(10.5)
 styles[name].paragraph_format.space_after=Pt(5)
 styles[name].paragraph_format.line_spacing=1.06
 styles[name].paragraph_format.widow_control=True
core=doc.core_properties
core.title='EKalusugan User Manual'
core.subject='Instructions for Patient Dentist Staff and Administrator roles'
core.author='E-KALUSUGAN'
core.keywords='EKalusugan user manual patient dentist staff administrator'

def clean(s):
 return str(s).replace('—',' - ').replace('–','-').replace('’',"'").replace('“','"').replace('”','"').replace('✓','').replace('←','').replace('·',' | ').strip()
def para(text,style=None,boldlead=None):
 pp=doc.add_paragraph(style=style)
 if boldlead and text.startswith(boldlead):
  pp.add_run(boldlead).bold=True; pp.add_run(clean(text[len(boldlead):]))
 else: pp.add_run(clean(text))
 return pp
break_pending=False
def heading(text,level=2):
 global break_pending
 pp=doc.add_heading(re.sub('[^A-Za-z0-9 ]',' ',text).replace('  ',' ').strip(),level)
 if break_pending:pp.paragraph_format.page_break_before=True;break_pending=False
 return pp
def page():
 global break_pending
 break_pending=True
def note(text):
 pp=doc.add_paragraph(); pp.add_run('Note  ').bold=True; pp.add_run(clean(text)); return pp
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
 pr.append(borders)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  for c,txt in zip(t.add_row().cells,row): c.text=clean(txt)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    v=OxmlElement('w:'+side);v.set(qn('w:w'),'80');v.set(qn('w:type'),'dxa');mar.append(v)
   tcPr.append(mar)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F6F8FA' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1.0
    for r in pp.runs:r.font.size=Pt(9.3);r.bold=(ri==0)
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

fig=0; figures=[]
def shot(id,title,maxh=3.45,crop=None):
 global fig
 path=SHOTS/(id+'.png')
 if not path.exists(): raise FileNotFoundError(path)
 meta=json.loads((SHOTS/(id+'.json')).read_text(encoding='utf-8'))
 if crop:
  im=Image.open(path); im=im.crop(crop); path=ROOT/'work'/(id+'_crop.png');im.save(path)
 im=Image.open(path);iw,ih=im.size;h=min(maxh,W*ih/iw);w=h*iw/ih
 pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.keep_with_next=True
 shape=pp.add_run().add_picture(str(path),width=Inches(w))
 shape._inline.docPr.set('descr',title+' with numbered controls described below')
 fig+=1;figures.append((fig,id,title))
 label=f'Figure {fig}  {title}'
 if meta.get('scenario'):label+='  Example records'
 para(label,'Caption')
 pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(8)
 pp.paragraph_format.line_spacing=1.0
 for i,l in enumerate(meta.get('labels',[])):
  if i:pp.add_run('   ')
  r=pp.add_run(str(l['number'])+'  ');r.bold=True;r.font.size=Pt(9)
  r=pp.add_run(clean(l['text']).rstrip('. ')+'.');r.font.size=Pt(9)

def section(s,title=None,notes=True):
 if title:heading(title,3)
 intro=s.get('intro') or s.get('purpose')
 if intro:para(intro)
 for i,x in enumerate(s.get('steps',[]),1):
  pp=doc.add_paragraph(); pp.paragraph_format.left_indent=Inches(.23);pp.paragraph_format.first_line_indent=Inches(-.23)
  pp.add_run(f'{i}. ').bold=True;pp.add_run(clean(x))
 for x in s.get('items',[]):para(x,'List Bullet')
 if 'status_table' in s:table(['Ticket display','What it means','What you should do'],[[x['status'],x['meaning'],x['action']] for x in s['status_table']],[1.05,2.16,3.76])
 if 'clinical_measures' in s:
  heading('Clinical measures in the entry form',3)
  for x in s['clinical_measures']:para(x,'List Bullet')
 if notes:
  for x in s.get('notes',[]): note(x)
def screen(title,id,sections,maxh=3.45):
 page(); heading(title,2);shot(id,title,maxh)
 for item in sections:
  if isinstance(item,tuple):section(item[0],item[1])
  else:section(item)

# Front matter
para('EKalusugan User Manual','Title')
para('Patient Dentist Staff and Administrator instructions','Subtitle')
para('Version 1.0   |   18 October 2026')
para('Use this manual to sign in, understand your dashboard, and complete the tasks assigned to your role. It covers the current dental queue, walk-in registration, patient records, clinical documentation, reports, account management, and SMS notifications.')
shot('public_home','EKalusugan homepage',3.8)
para('Read the access instructions first, then follow the chapter for your role. Red numbered labels identify the controls explained below each screenshot. Names, contact details, and account identities in the screenshots are anonymized. Example queue and clinical records illustrate actions in the actual running interface.')
para('The patient portal currently handles same-day dental queue tickets and appointment history. Clinic administrators provide accounts; public self-registration and future-date patient booking are not available in this version.')

page();heading('Using this manual',1)
para('Before you start, obtain the EKalusugan website address and an active account from the clinic. Use a browser with an internet connection. Keep your registered mobile phone available if the clinic sends your first temporary password by SMS.')
para('A button name in these instructions matches the system text. A disabled button appears faded and cannot be used until its requirements are met. Read any message near the button before retrying. Close or Cancel in a dialog generally returns to the previous screen; Continue or a save button confirms the described action.')
heading('Choose your role',2)
table(['Role','Main responsibilities','Begin with'],[
 ['Patient','Request and follow your own dental ticket, review appointment history, and find the clinic.','Chapter 2'],
 ['Dentist or Doctor','Set availability and daily limits, manage the dental queue, save consultation notes, enter clinical counts, and export reports.','Chapter 4'],
 ['Staff','Find patients, register dental walk-ins, manage arrival and queue actions, review records, and verify in-person recovery.','Chapter 3'],
 ['Administrator','Oversee queues and activity, manage accounts, view history and forecasts, and configure SMS alerts.','Chapter 5']],[1.26,4.55,1.16])
heading('Read the screen correctly',2)
para('The system uses Doctor as the account role and Dentist for the dental service. A ticket number identifies a queue visit. An appointment record and a clinical report are separate records. A blank, dash, or loading indicator does not confirm that a count is zero.')
para('Dates and queue return times use Philippine time. Patient and Staff screens refresh about every 15 seconds; Dentist screens refresh about every 30 seconds. Read the latest state before acting because another authorized user may have updated the same ticket.')
heading('Protect patient information',2)
para('Use your own account. Verify the selected patient before opening or changing records. Keep passwords private, share a recovery password only with the verified patient, and use Log Out when finished. Keep the public queue monitor on the dedicated display and keep account and clinical screens at the staff workstation.')

page();heading('Contents',1)
pp=doc.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');pp._p.append(fld)

page();heading('1 Access and sign in',1)
heading('Open the homepage',2)
shot('public_home','Homepage navigation')
for x in ['Open the EKalusugan address supplied by the clinic.','Select Log in in the upper right. Why E-Kalusugan scrolls to the feature overview; Your next visit scrolls to the invitation to sign in.','The homepage preview is an introduction to the portal. Sign in with your clinic-issued account to reach your role dashboard.']:
 para(x,'List Number')
note('There is no active public registration page. Contact clinic reception if you do not yet have an account.')
screen('Sign in to your account','public_login',[{'steps':[
 'Enter your registered mobile number in Phone Number. Use the Philippine mobile format shown in the field, such as 09 followed by nine digits.',
 'Enter your password. Use the eye icon only when it is safe to display the password on screen.',
 'Select Login once and wait while Logging in is displayed.',
 'A Patient, Doctor, Staff, or Administrator account opens its assigned dashboard automatically. You do not choose a role on the login form.',
 'If the system requires a new password, complete Change Your Password before continuing.'
 ],'notes':['If Please fill in all fields appears, complete both fields. If Invalid phone or password appears, check the registered phone and password. Ask reception for recovery assistance if access still fails.']}])
screen('Recover access through reception','public_recovery',[{'steps':[
 'On the login form, select Forgot password? to expand Password Recovery Steps.',
 'Visit clinic reception. Verify your name, registered login phone number, address, and birthday with authorized staff.',
 'Receive the temporary recovery password privately. It is valid for 30 minutes and one login.',
 'Sign in with the same registered phone and the recovery password. Set a permanent password when prompted.'
 ],'notes':['Forgot password? displays instructions for in-person recovery. It does not send a reset link or change your password.','The first activation password sent by SMS is a different workflow. Follow the activation message and complete password setup promptly; ask the clinic if it expires or cannot be used.']}])
screen('Set or change your password','password_change',[{'steps':[
 'At first activation or recovery login, the system opens Change Your Password. In this setup flow, the form asks for New Password and Confirm New Password.',
 'For a normal password change, enter Current Password as well, as shown in the screenshot.',
 'Create a password with at least eight characters and at least one number. The maximum is 72 bytes, equivalent to 72 plain English characters and fewer for accented characters or emoji.',
 'Enter the same new password in Confirm New Password. Use the eye controls to check the fields when appropriate.',
 'Select Update Password. Wait for the assigned dashboard to open. If Passwords do not match appears, enter matching values and try again.'
 ],'notes':['A temporary-password setup session is time limited. If it expires, return to reception or request the appropriate activation assistance.','You can open the normal password-change page from the website address ending in /change-password while signed in.']}])
if (SHOTS/'public_terms.png').exists():
 screen('Read terms and end your session','public_terms',[{'steps':[
 'Select terms and conditions or privacy policy on the login page to open Terms, Privacy & Consent.',
 'Read the policy and scroll through the policy area. The consent controls become available after reaching the bottom.',
 'Select the consent checkbox if you agree, then select Accept & Continue. Use Cancel or the X to close the dialog when you only need to return to sign in.',
 'When finished with a dashboard, open your name or profile chip at the upper right and select Log Out. On a phone, open the menu first if the profile control is hidden.'
 ],'notes':['Closing the browser tab ends that tab session. Sign in again if you return to the login page or see an expired-session message.']}])

page();heading('2 Patient instructions',1)
para(ps['role_boundaries'][0]['purpose'])
heading('Understand the patient dashboard',2);shot('patient_home','Patient dashboard');section(p['patient-dashboard'])
screen('Request a dental queue ticket','patient_ticket',[p['patient-get-ticket']])
page();heading('Dental service choices',2)
para('Choose one or two services for one dental visit. Both the patient ticket form and staff walk-in form use these choices.')
table(['Service group','Choices shown in the forms'],[
 ['General dental care','Dental Check-up; Oral Prophylaxis; Fluoride; Silver Diamine; Special Surgery; Others'],
 ['Restoration or tooth filling','Permanent Filling; Temporary Filling'],
 ['Removable partial denture','Removable Partial Denture - Upper; Removable Partial Denture - Lower'],
 ['Dental surgery','Closed Extraction; Open Extraction; Odontectomy / Wisdom Tooth Extraction']],[1.9,5.07])
heading('Read your active ticket',2);shot('patient_active','Active patient ticket',3.0);section({'steps':[
 'Read Your Queue Number and the service labels. Use the ticket number when reporting to staff.',
 'Read the status and the instruction above the ticket. If no SMS will be sent, stay nearby. If one alert is expected, also follow the queue display.',
 'When called, report to staff immediately so they can confirm your presence.'
 ],'notes':['Queue order may change. Initial position is not a fixed place, and the displayed estimated wait is indicative.']})
page();heading('Ticket statuses and the next action',2);section(p['patient-status'])
screen('Cancel a waiting or missed ticket','patient_cancel',[p['patient-cancel']])
screen('Return after a missed call','patient_missed',[{'steps':[
 'Read the deadline shown above the ticket after Missed call appears.',
 'Report to Staff before that exact time. Staff must select Patient returned to confirm your return.',
 'Stay nearby after confirmation. The queue continues while you return, and the restored ticket is called at the next opening.',
 'If the deadline expires or a returning patient misses the next call, the ticket ends as No-show. A new request depends on current availability and daily limits.'
 ],'notes':['Refreshing the patient screen does not confirm a return. Do not wait for an SMS to report to reception.']}])
screen('Review history and find the clinic','patient_history',[p['patient-history-location']])

page();heading('3 Staff instructions',1)
para(ps['role_boundaries'][1]['purpose'])
heading('Understand the staff dashboard',2);shot('staff_home','Staff dashboard');section(p['staff-dashboard'])
screen('Find an existing patient','staff_find',[p['staff-find-patient']],3.5)
screen('Register a dental walk in','staff_services',[p['staff-register-walkin']],3.7)
screen('Call and serve dental patients','staff_called',[p['staff-serve']],3.7)
page();heading('Handle missed calls and returning patients',2);section(p['staff-missed-return'])
heading('Cancel a ticket with a reason',2);section(p['staff-cancel'])
screen('Use the dental queue monitor','staff_monitor',[{'steps':[
 'From the staff dashboard, select Open dental monitor. The monitor opens in another browser tab.',
 'Keep this display visible to waiting patients. Read the current and next dental queue numbers and the waiting list.',
 'Select Enable announcements using the speaker control if audio announcements are needed, and allow audio through the browser.',
 'Keep the staff dashboard available in its own tab for calling, confirming arrival, and handling returns.'
 ],'notes':['The monitor requires a Staff or Administrator session. If the tab returns to sign in, sign in with an authorized account and reopen it.']}])
screen('Review the patient profile','staff_profile',[p['staff-profile']])
screen('Merge a duplicate walk in record','staff_merge',[p['staff-merge']])
screen('Find a patient for account recovery','staff_recovery',[{'steps':p['staff-account-recovery']['steps'][:3],'notes':['Account recovery is for a patient who is physically present. Search for the correct record and open the profile before starting the verification checklist.']}])
screen('Verify identity and issue recovery access','staff_recovery_verify',[{'intro':p['staff-account-recovery']['intro'],'steps':p['staff-account-recovery']['steps'][3:],'notes':p['staff-account-recovery']['notes']}])

page();heading('4 Dentist instructions',1)
para('Manage today\'s dental service, document consultations, and prepare clinical reports from the Doctor account.')
heading('Understand the dentist dashboard',2);shot('doctor_home','Dentist dashboard',3.3);section(d[0],notes=False)
screen('Set availability and daily capacity','doctor_capacity',[(d[1],'Set availability'),(d[2],'Set patient limits')])
screen('Call patients and confirm arrival','doctor_called',[d[3]])
screen('Document and complete the consultation','doctor_consultation',[d[4]])
screen('Review appointment history','doctor_history',[d[5]])
screen('Read dentist analytics','doctor_analytics',[d[6]])
screen('Enter daily barangay clinical counts','doctor_report_entry',[{'intro':d[7]['purpose'],'steps':d[7]['steps'][:4]}],4.3)
page();heading('Save and check clinical counts',2)
section({'steps':d[7]['steps'][4:],'clinical_measures':d[7]['clinical_measures'],'notes':d[7]['notes']})
screen('Print or download reports','doctor_report',[d[8]])

page();heading('5 Administrator instructions',1)
para('Oversee service activity, manage accounts, review patient history, and configure SMS notifications.')
heading('Navigate the administrator dashboard',2);shot('admin_home','Administrator dashboard');section(a[0]);section(a[1],'Review activity and volume',notes=False)
screen('Monitor and cancel queue tickets','admin_queue',[a[2]],1.75)
if (SHOTS/'admin_cancel.png').exists():
 heading('Choose a cancellation reason',2);shot('admin_cancel','Administrator cancellation reason',2.6)
 para('Select one reason and then Continue to confirm. Select Cancel to keep the ticket. The reason is stored with the cancelled visit.')
screen('Review statistics and forecasts','admin_statistics',[a[3]])
screen('Check doctor availability','admin_doctors',[a[4]])
screen('Review patient appointment history','admin_history',[a[5]])
screen('Configure and monitor SMS notifications','admin_notifications',[a[6]])
screen('Manage staff doctor and admin accounts','admin_staff',[{'steps':[
 'Select Accounts in the sidebar, then select Staff. The Staff Management list includes staff, doctor, and administrator accounts.',
 'Use Name, Email, Role, and Status in the list to identify the correct person before choosing an action.',
 'Select Add Staff to open a new account form. Use Edit on an existing account to update it.'
 ]},(a[8],'Edit and manage account status')])
screen('Create a staff doctor or admin account','admin_staff_form',[a[7]])
screen('Find and manage patient accounts','admin_patients',[{'steps':[
 'Select Accounts, then Patients to open Patient Accounts.',
 'Find the patient using the search and filters. Check account state, activation or password SMS status, and the matching identity.',
 'Use Previous and Next to move through pages. Select Add Patient only after checking for an existing account.'
 ]},(a[10],'Manage activation and access')])
screen('Create a patient account','admin_patient_form',[a[9]])
page();heading('Assist a verified patient with recovery',2);section(a[11])
para('Follow the illustrated reception checklist in Chapter 3 for the verification screen. Activation SMS replacement and in-person account recovery have different actions; select the workflow that matches the patient\'s account state.')

page();heading('6 Common problems and routine checks',1)
heading('Sign in and account access',2)
table(['What you see','What to do'],[
 ['Please fill in all fields','Enter the registered phone number and password.'],
 ['Invalid phone or password','Check the number and password. Use Forgot password? for reception recovery instructions.'],
 ['Temporary password fails or expires','Ask the clinic to check activation status or complete in-person recovery. Keep the same registered login phone.'],
 ['Change Your Password appears','Complete the required password setup before using a dashboard.'],
 ['Session expired or login page appears again','Sign in again. If access remains unavailable, ask the administrator to check that your account is active.'],
 ['Wrong dashboard for the account','Open your profile and check the displayed role. Ask the administrator to check the account assignment.']],[2.2,4.77])
heading('Queue and walk in problems',2)
table(['What you see','What to do'],[
 ['Get queue ticket is disabled','Read the nearby message. An active ticket, registration closure, or daily limit may prevent another request.'],
 ['An old number or delayed status','Wait for the automatic refresh. Staff can use Refresh. Check the connection before repeating an action.'],
 ['Call again or Skip is disabled','Allow the displayed 30-second wait to finish. The second call must occur before skipping.'],
 ['Missed call','Report to staff before the displayed return deadline. Staff confirms the return.'],
 ['Possible patient match found','Review the existing records. Select the correct patient, or confirm a new record only after checking identity details.'],
 ['Mobile number is rejected','Use the requested format. Walk-in entry expects ten digits starting with 9 after the displayed +63 prefix.'],
 ['No queue SMS arrives','Follow the on-screen ticket instruction. Some tickets intentionally receive no SMS; delivery can also be delayed.']],[2.2,4.77])

page();heading('Reports notifications and records',2)
table(['What you see','What to do'],[
 ['Report export is disabled','Save clinical entries, resolve load errors, and wait for loading to finish. The selected period needs completed visits or manual entries.'],
 ['Unsaved changes','Use Save entries before changing the date or period. Read any discard confirmation carefully.'],
 ['Dentist profile not found','Ask the administrator to check the Doctor account and linked dentist profile.'],
 ['SMS is queued or processing','Wait for the status to update. Queued does not mean delivered or read.'],
 ['SMS is unknown or failed','Review the activity status. Follow the clinic\'s provider checks before issuing a replacement password; do not repeatedly create accounts.'],
 ['Recovery button is unavailable','Check active account eligibility, complete identity details, and finish the required verification boxes.'],
 ['No appointments or no patients waiting','This is a valid empty state when no matching records exist. Check the selected day, period, and filters.']],[2.2,4.77])
heading('Daily routine by role',2)
for role,items in [
 ('Patient',['Check availability and registration before requesting a ticket.','Stay near the clinic when instructed and report promptly when called.','Review your appointment history and log out.']),
 ('Staff',['Check dentist availability and refresh the queue.','Find the existing record before creating a walk-in record.','Confirm patient contact, priority eligibility, services, and the registration SMS instruction.','Record arrival, missed returns, completion, and cancellation accurately.']),
 ('Dentist',['Set today\'s availability and patient limits.','Check the correct patient before entering and saving consultation notes.','Review and save verified clinical counts before printing or downloading reports.']),
 ('Administrator',['Check accounts and activation states.','Review queue activity and any SMS delivery problems.','Review the selected period and data limitations before using statistics.'])]:
 heading(role,3)
 for x in items:para(x,'List Bullet')
heading('Get help from the clinic',2)
para('For a problem that remains unresolved, note your role, the page name, the action you selected, and the exact error message. Give this information to clinic reception or the system administrator. Do not include a password in a support message.')

# Page numbering is useful in a multi-role reference manual.
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=foot.add_run('EKalusugan User Manual   |   ');run.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
settings=doc.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
doc.save(DOC)
(ROOT/'work/figure_inventory.json').write_text(json.dumps(figures,indent=2),encoding='utf-8')
print(f'Created {DOC} with {fig} figures')
